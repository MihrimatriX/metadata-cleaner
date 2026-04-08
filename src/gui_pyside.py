from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QLabel, QListWidget, QListWidgetItem, QFileDialog, QMessageBox, QAbstractItemView, QFrame, QTableWidget, QTableWidgetItem, QSplitter, QCheckBox
)
from PySide6.QtGui import QIcon, QColor, QBrush, QPixmap
from PySide6.QtCore import Qt, QTimer
import sys
sys.dont_write_bytecode = True
import os
import threading
from remover import remove_metadata, SUPPORTED_EXTENSIONS
from concurrent.futures import ThreadPoolExecutor

# Emoji icons by file type
FILE_ICONS = {
    # Image
    '.jpg': '🖼️', '.jpeg': '🖼️', '.png': '🖼️', '.tiff': '🖼️', '.bmp': '🖼️', '.webp': '🖼️', '.heic': '🖼️',
    # Document
    '.pdf': '📄', '.docx': '📄', '.doc': '📄', '.odt': '📄', '.epub': '📄', '.txt': '📄', '.rtf': '📄', '.csv': '📄', '.xlsx': '📊', '.xls': '📊', '.pptx': '📊', '.ppt': '📊', '.ods': '📊', '.odp': '📊',
    # Audio
    '.mp3': '🎵', '.wav': '🎵', '.flac': '🎵', '.ogg': '🎵', '.aac': '🎵', '.wma': '🎵', '.m4a': '🎵', '.aiff': '🎵',
    # Video
    '.mp4': '🎬', '.mkv': '🎬', '.mov': '🎬', '.avi': '🎬', '.wmv': '🎬', '.webm': '🎬', '.m4v': '🎬',
    # Archive
    '.zip': '🗜️', '.rar': '🗜️', '.7z': '🗜️', '.tar': '🗜️', '.gz': '🗜️',
    # Code
    '.py': '💻', '.js': '💻', '.html': '💻', '.css': '💻', '.json': '💻', '.xml': '💻',
    'default': '📁'
}

def get_icon_for_file(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    return FILE_ICONS.get(ext, FILE_ICONS['default'])

# ModernBar and ModernFileListItem removed, reverted to old FileListItem
class FileListItem(QWidget):
    def __init__(self, filename, status=None, on_remove=None):
        super().__init__()
        layout = QHBoxLayout()
        layout.setContentsMargins(4, 2, 4, 2)
        layout.setSpacing(8)
        # Silme butonu
        self.remove_btn = QPushButton('×')
        self.remove_btn.setFixedWidth(24)
        self.remove_btn.setStyleSheet('color:#e57373; font-weight:bold; background:transparent; border:none;')
        if on_remove:
            self.remove_btn.clicked.connect(on_remove)
        layout.addWidget(self.remove_btn)
        # Icon
        icon_label = QLabel(get_icon_for_file(filename))
        icon_label.setFixedWidth(28)
        icon_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(icon_label)
        # File name
        name_label = QLabel(os.path.basename(filename))
        name_label.setMinimumWidth(120)
        layout.addWidget(name_label)
        # Durum
        self.status_label = QLabel(status or '')
        self.status_label.setFixedWidth(32)
        self.status_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(self.status_label)
        self.setLayout(layout)

class MetadataCleanerGUI(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Metadata Cleaner")
        self.setGeometry(200, 200, 800, 540)
        self.setAcceptDrops(True)
        self.file_list = []  # (filename, status)
        self.executor = ThreadPoolExecutor(max_workers=4)
        self._clean_lock = threading.Lock()
        self.init_ui()

    def init_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        main_layout = QVBoxLayout()
        main_layout.setContentsMargins(0,0,0,0)
        main_layout.setSpacing(0)
        central.setLayout(main_layout)

        # Top bar
        top_bar = QHBoxLayout()
        add_file_btn = QPushButton("Add File")
        add_file_btn.clicked.connect(self.add_file)
        add_folder_btn = QPushButton("Add Folder")
        add_folder_btn.clicked.connect(self.add_folder)
        self.chk_subfolders = QCheckBox("Include subfolders")
        self.chk_subfolders.setToolTip("When adding a folder, also scan nested directories.")
        top_bar.addWidget(add_file_btn)
        top_bar.addWidget(add_folder_btn)
        top_bar.addWidget(self.chk_subfolders)
        top_bar.addStretch()
        main_layout.addLayout(top_bar)
        # Splitter with main list and detail panel
        self.splitter = QSplitter()
        self.splitter.setOrientation(Qt.Horizontal)
        main_layout.addWidget(self.splitter)
        # Left: File list panel
        self.list_panel = QWidget()
        list_layout = QVBoxLayout()
        self.list_panel.setLayout(list_layout)
        self.list_widget = QListWidget()
        self.list_widget.setSelectionMode(QAbstractItemView.SingleSelection)
        self.list_widget.setFrameShape(QFrame.NoFrame)
        self.list_widget.itemSelectionChanged.connect(self.handle_selection_changed)
        list_layout.addWidget(self.list_widget)
        # Bottom bar
        bottom_bar = QHBoxLayout()
        self.status_label = QLabel("No files added yet.")
        bottom_bar.addWidget(self.status_label)
        bottom_bar.addStretch()
        self.clean_btn = QPushButton("Clean")
        self.clean_btn.setStyleSheet("background-color: #e57373; color: white; font-weight: bold;")
        self.clean_btn.clicked.connect(self.clean_files)
        bottom_bar.addWidget(self.clean_btn)
        list_layout.addLayout(bottom_bar)
        self.splitter.addWidget(self.list_panel)
        # Right: Detail panel (always visible)
        self.detail_panel = QWidget()
        detail_layout = QVBoxLayout()
        self.detail_panel.setLayout(detail_layout)
        self.detail_title = QLabel("Details")
        self.detail_title.setStyleSheet("font-weight:bold;font-size:16px;")
        detail_layout.addWidget(self.detail_title)
        self.detail_table = QTableWidget()
        self.detail_table.setColumnCount(2)
        self.detail_table.setHorizontalHeaderLabels(["Field", "Value"])
        self.detail_table.horizontalHeader().setStretchLastSection(True)
        detail_layout.addWidget(self.detail_table)
        self.splitter.addWidget(self.detail_panel)

    def handle_selection_changed(self):
        selected = self.list_widget.selectedItems()
        if selected:
            idx = self.list_widget.row(selected[0])
            self.show_details(idx)
        else:
            self.clear_details()

    def clear_details(self):
        self.detail_title.setText("")
        self.detail_table.setRowCount(0)

    def add_file(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "Select File")
        if file_path:
            self.file_list.append({'filename': file_path, 'status': ''})
            self.refresh_list()

    def _collect_from_folder(self, folder_path: str, recursive: bool) -> None:
        exts = SUPPORTED_EXTENSIONS.keys()
        if recursive:
            for root, _, files in os.walk(folder_path):
                for f in files:
                    if os.path.splitext(f)[1].lower() in exts:
                        self.file_list.append({'filename': os.path.join(root, f), 'status': ''})
        else:
            for fname in os.listdir(folder_path):
                fpath = os.path.join(folder_path, fname)
                if os.path.isfile(fpath) and os.path.splitext(fname)[1].lower() in exts:
                    self.file_list.append({'filename': fpath, 'status': ''})

    def add_folder(self):
        folder_path = QFileDialog.getExistingDirectory(self, "Select Folder")
        if folder_path:
            self._collect_from_folder(folder_path, self.chk_subfolders.isChecked())
            self.refresh_list()

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            super().dragEnterEvent(event)

    def dropEvent(self, event):
        for url in event.mimeData().urls():
            path = url.toLocalFile()
            if not path:
                continue
            if os.path.isfile(path):
                if os.path.splitext(path)[1].lower() in SUPPORTED_EXTENSIONS:
                    self.file_list.append({'filename': path, 'status': ''})
            elif os.path.isdir(path):
                self._collect_from_folder(path, self.chk_subfolders.isChecked())
        self.refresh_list()
        event.acceptProposedAction()

    def remove_file(self, idx):
        del self.file_list[idx]
        self.refresh_list()

    def show_details(self, idx):
        file_path = self.file_list[idx]['filename']
        self.detail_title.setText(os.path.basename(file_path))
        meta = self.get_metadata(file_path)
        self.detail_table.setRowCount(len(meta))
        for i, (k, v) in enumerate(meta.items()):
            self.detail_table.setItem(i, 0, QTableWidgetItem(str(k)))
            self.detail_table.setItem(i, 1, QTableWidgetItem(str(v)))
        self.detail_panel.show()
        self.splitter.setSizes([300, 500])

    def hide_details(self):
        self.detail_panel.hide()
        self.splitter.setSizes([1, 0])

    def get_metadata(self, file_path):
        ext = os.path.splitext(file_path)[1].lower()
        try:
            if ext in ['.jpg', '.jpeg', '.png', '.tiff', '.bmp', '.webp', '.heic']:
                from PIL import Image
                img = Image.open(file_path)
                exif = img.getexif()
                return {k: v for k, v in exif.items()}
            elif ext in ['.pdf']:
                from pypdf import PdfReader
                reader = PdfReader(file_path)
                return reader.metadata or {}
            elif ext in ['.docx', '.doc', '.odt', '.rtf']:
                from docx import Document
                doc = Document(file_path)
                props = doc.core_properties
                return {k: getattr(props, k) for k in dir(props) if not k.startswith('_') and not callable(getattr(props, k))}
            elif ext in ['.xlsx', '.xls', '.ods']:
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(file_path, read_only=True)
                    props = wb.properties
                    return {k: getattr(props, k) for k in dir(props) if not k.startswith('_') and not callable(getattr(props, k))}
                except Exception:
                    return {'Info': 'Excel file, showing basic info.'}
            elif ext in ['.pptx', '.ppt', '.odp']:
                try:
                    from pptx import Presentation
                    prs = Presentation(file_path)
                    props = prs.core_properties
                    return {k: getattr(props, k) for k in dir(props) if not k.startswith('_') and not callable(getattr(props, k))}
                except Exception:
                    return {'Info': 'Presentation file, showing basic info.'}
            elif ext in ['.txt', '.csv', '.json', '.xml', '.html', '.css', '.js', '.py']:
                size = os.path.getsize(file_path)
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    lines = f.readlines()
                return {'Size (bytes)': size, 'Line count': len(lines)}
            elif ext in ['.mp3', '.flac', '.ogg', '.wav', '.aac', '.wma', '.m4a', '.aiff']:
                from mutagen import File
                audio = File(file_path)
                return dict(audio.tags) if audio and audio.tags else {}
            elif ext in ['.mp4', '.mkv', '.avi', '.mov', '.wmv', '.webm', '.m4v']:
                from pymediainfo import MediaInfo
                info = MediaInfo.parse(file_path)
                return {t.track_type: t.to_data() for t in info.tracks}
            elif ext in ['.zip', '.rar', '.7z', '.tar', '.gz']:
                size = os.path.getsize(file_path)
                return {'Size (bytes)': size}
            else:
                return {"Info": "No details available for this file type."}
        except Exception as e:
            return {"Error": str(e)}

    def refresh_list(self):
        self.list_widget.clear()
        for idx, item in enumerate(self.file_list):
            widget = FileListItem(
                item['filename'],
                status=item.get('status', ''),
                on_remove=self.make_remove_handler(idx)
            )
            list_item = QListWidgetItem()
            list_item.setSizeHint(widget.sizeHint())
            self.list_widget.addItem(list_item)
            self.list_widget.setItemWidget(list_item, widget)
        self.status_label.setText(f"{len(self.file_list)} files added.")
        # Auto-select: select first file if exists
        if self.file_list:
            self.list_widget.setCurrentRow(0)
        else:
            self.clear_details()

    def make_remove_handler(self, idx):
        return lambda checked=False, i=idx: self.remove_file(i)

    def clean_files(self):
        if not self.file_list:
            QMessageBox.warning(self, "Warning", "Please add a file or folder first.")
            return
        self.clean_btn.setEnabled(False)
        self.status_label.setText("Operation started...")
        self.cleaned_count = 0
        self.failed_count = 0
        self._completed = 0
        self._total_to_clean = len(self.file_list)
        for idx, item in enumerate(self.file_list):
            self.executor.submit(self.clean_single_file, idx, item)

    def clean_single_file(self, idx, item):
        trigger_done = False
        try:
            result = remove_metadata(item['filename'])
            with self._clean_lock:
                if result:
                    item['status'] = '✓'
                    self.cleaned_count += 1
                else:
                    item['status'] = '✗'
                    self.failed_count += 1
                self._completed += 1
                trigger_done = self._completed == self._total_to_clean
        except Exception:
            with self._clean_lock:
                item['status'] = '✗'
                self.failed_count += 1
                self._completed += 1
                trigger_done = self._completed == self._total_to_clean
        QTimer.singleShot(0, self.refresh_list)
        if trigger_done:
            QTimer.singleShot(0, self.cleaning_done)

    def cleaning_done(self):
        self.clean_btn.setEnabled(True)
        self.status_label.setText(f"{self.cleaned_count} files cleaned, {self.failed_count} errors.")
        QMessageBox.information(self, "Operation completed", f"{self.cleaned_count} files cleaned successfully. {self.failed_count} errors.")

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MetadataCleanerGUI()
    window.show()
    sys.exit(app.exec()) 